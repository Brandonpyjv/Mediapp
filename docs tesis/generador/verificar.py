"""
Comprobador de estilo de los entregables.

Revisa un `.docx` ya generado y señala lo que el §1.1 del cuaderno prohíbe, más la
regla que el evaluador escribió sobre los casos de uso. No comprueba el formato,
de eso se encarga `apa.py`, que lo aplica por código y no depende de acordarse.

    python verificar.py ../entregables/archivo.docx
    python verificar.py ../entregables/*.docx

**Por qué existe desde T2 y no desde el ensamble final.** En FactuGest estas
mismas reglas se comprobaron al terminar, y hubo que corregir los cuatro
entregables completos de una sentada. Teniéndolo desde el principio, cada tarea
revisa lo suyo el día que lo escribe, que es cuando corregirlo cuesta un minuto.
"""
import re
import sys
from pathlib import Path

from docx import Document

# El guion largo y el guion medio. Casi nadie los escribe a mano, y el autor de
# FactuGest los detectó como marca de texto generado.
GUIONES = re.compile(r"[—–]")

# El punto medio, que cayó en la misma revisión y por la misma razón.
PUNTO_MEDIO = re.compile(r"·")

# El patrón «afirmación breve: explicación», la tercera marca. Se señala para
# revisar y no como falta segura, porque hay tres usos legítimos: la entrada de
# glosario con viñeta, la etiqueta de dato («Palabras clave:») y el título real de
# una obra citada. Quien revisa decide.
DOS_PUNTOS = re.compile(r"[a-záéíóúñ]{4,}: [a-záéíóúñ][a-záéíóúñ]+")

# O12: «En los casos de uso no debe ir codigo SQL».
SQL = re.compile(r"\b(SELECT|INSERT\s+INTO|UPDATE|DELETE\s+FROM|CREATE\s+TABLE|"
                 r"ALTER\s+TABLE|JOIN|WHERE)\b", re.IGNORECASE)


def _textos(ruta):
    """Todo el texto del documento, incluido el de las tablas.

    Las celdas se recorren aparte porque `document.paragraphs` no entra en ellas,
    y es justo donde se colaban los guiones largos de FactuGest, que los usaba
    como marcador de celda vacía.
    """
    documento = Document(str(ruta))
    for numero, parrafo in enumerate(documento.paragraphs, 1):
        if parrafo.text.strip():
            yield f"párrafo {numero}", parrafo.text
    for indice, tabla in enumerate(documento.tables, 1):
        for fila in tabla.rows:
            for celda in fila.cells:
                if celda.text.strip():
                    yield f"tabla {indice}", celda.text


def _fragmento(texto, posicion, margen=45):
    inicio = max(0, posicion - margen)
    return ("..." if inicio else "") + texto[inicio:posicion + margen].replace("\n", " ") + "..."


def revisar(ruta, con_sql=False):
    """Devuelve (faltas, avisos). `con_sql` solo aplica a los casos de uso."""
    faltas, avisos = [], []
    reglas = [("guion largo", GUIONES), ("punto medio", PUNTO_MEDIO)]
    if con_sql:
        reglas.append(("código SQL", SQL))

    for donde, texto in _textos(ruta):
        for nombre, patron in reglas:
            for encontrado in patron.finditer(texto):
                faltas.append((nombre, donde, _fragmento(texto, encontrado.start())))
        for encontrado in DOS_PUNTOS.finditer(texto):
            avisos.append(("«afirmación: explicación»", donde,
                           _fragmento(texto, encontrado.start())))
    return faltas, avisos


def informar(ruta, con_sql=False):
    faltas, avisos = revisar(ruta, con_sql)
    print(f"\n{Path(ruta).name}")
    if not faltas and not avisos:
        print("   sin observaciones")
    for nombre, donde, fragmento in faltas:
        print(f"   FALTA  [{nombre}] {donde}: {fragmento}")
    for nombre, donde, fragmento in avisos:
        print(f"   revisar [{nombre}] {donde}: {fragmento}")
    print(f"   {len(faltas)} faltas, {len(avisos)} por revisar")
    return len(faltas)


if __name__ == "__main__":
    archivos = sys.argv[1:]
    if not archivos:
        print(__doc__)
        sys.exit(1)
    # Los casos de uso son los únicos donde el SQL es falta; en el resto del
    # documento una consulta de ejemplo puede tener su sitio.
    total = sum(informar(a, con_sql="Casos de Uso" in a) for a in archivos)
    sys.exit(1 if total else 0)
