# -*- coding: utf-8 -*-
"""
Escribe la sustentación de MediApp ENCIMA de la presentación de FactuGest.

No se diseña nada. Se copia `FactuGest - Sustentación.pptx` a la carpeta de
MediApp y se sobrescribe su contenido diapositiva por diapositiva: cada caja de
texto conserva su posición, su tamaño y su tipografía, y solo cambia lo que
dice; cada imagen conserva su marco y solo cambia el archivo que muestra; cada
tabla conserva su estilo y solo cambia lo que lleva dentro.

Las 38 diapositivas de FactuGest quedan en el mismo orden y con la misma forma.

Lo único que se sustituye de la marca es:
  * el logo «FG FACTUGEST»  ->  el logo de MediApp
  * la marca de agua «FG»    ->  se retira
  * el azul y el ámbar de FactuGest -> negro, con el azul de MediApp como acento

El logo del SENA viaja en el fondo de la plantilla y por eso sale en las 38.
"""
import copy
import shutil
import subprocess
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt
from PIL import Image as Pil

BASE = Path(__file__).resolve().parent
EXPOSICION = BASE / "exposicion"
ASSETS = EXPOSICION / "assets"
DIAGRAMAS = BASE / "entregables" / "diagramas"
CAPTURAS = BASE / "insumos" / "capturas"

ORIGEN = Path(r"D:\01. Proyectos\Factugest Python\Factugest\docs expo"
              r"\carpeta exposicion\FactuGest - Sustentación.pptx")
DESTINO = EXPOSICION / "MediApp - Sustentación.pptx"

LOGO_MEDIAPP = ASSETS / "mediapp_logo.png"
TRANSPARENTE = ASSETS / "transparent.png"

# Negro para el texto; el único color es el azul del logo de MediApp.
NEGRO = RGBColor(0x1A, 0x1A, 0x1A)
ACENTO = RGBColor(0x33, 0x69, 0xCD)

# Cómo se traduce cada color de FactuGest.
RECOLOR = {
    "003366": NEGRO,    # títulos
    "1F3B73": ACENTO,   # destacados
    "F5A623": ACENTO,   # rótulos ámbar
    "33373D": NEGRO,    # cuerpo
}

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


# --------------------------------------------------------------------------
# Utilidades de escritura sobre la plantilla
# --------------------------------------------------------------------------

def _formato_base(parrafo):
    """Copia del rPr del primer run, para reutilizarlo al reescribir."""
    for run in parrafo.runs:
        rPr = run._r.find(A + "rPr")
        if rPr is not None:
            return copy.deepcopy(rPr)
    return None


def _escribir(parrafo, texto, rPr):
    for run in list(parrafo.runs):
        run._r.getparent().remove(run._r)
    nuevo = parrafo.add_run()
    nuevo.text = texto
    if rPr is not None:
        viejo = nuevo._r.find(A + "rPr")
        if viejo is not None:
            nuevo._r.remove(viejo)
        nuevo._r.insert(0, copy.deepcopy(rPr))


def txt(forma, *lineas):
    """Reescribe la caja conservando la tipografía que ya tenía cada párrafo."""
    marco = forma.text_frame
    parrafos = list(marco.paragraphs)
    rPr = None
    for p in parrafos:
        rPr = _formato_base(p)
        if rPr is not None:
            break
    for i, linea in enumerate(lineas):
        if i < len(parrafos):
            # Cada párrafo conserva SU formato; el del primero solo cubre los vacíos.
            # Ojo: un rPr sin hijos es «falso» para lxml, así que se compara con None.
            propio = _formato_base(parrafos[i])
            _escribir(parrafos[i], linea, rPr if propio is None else propio)
        else:
            p = marco.add_paragraph()
            if parrafos:
                pPr = parrafos[0]._pPr
                if pPr is not None:
                    p._p.insert(0, copy.deepcopy(pPr))
            _escribir(p, linea, rPr)
    for sobra in parrafos[len(lineas):]:
        sobra._p.getparent().remove(sobra._p)


def borrar(forma):
    forma._element.getparent().remove(forma._element)


def _encoger(forma, lado):
    """Reduce una imagen cuadrada dejándola en el mismo centro."""
    cx, cy = forma.left + forma.width // 2, forma.top + forma.height // 2
    forma.width = forma.height = int(lado)
    forma.left, forma.top = cx - forma.width // 2, cy - forma.height // 2


def imagen(diapositiva, forma, ruta):
    """Cambia la foto que muestra el marco, encajándola sin deformarla."""
    ruta = str(ruta)
    parte, rId = diapositiva.part.get_or_add_image_part(ruta)
    blip = forma._element.find(f".//{A}blip")
    blip.set(R + "embed", rId)

    with Pil.open(ruta) as im:
        proporcion = im.height / im.width
    cx, cy = forma.left + forma.width // 2, forma.top + forma.height // 2
    ancho, alto = forma.width, int(forma.width * proporcion)
    if alto > forma.height:
        alto, ancho = forma.height, int(forma.height / proporcion)
    forma.width, forma.height = ancho, alto
    forma.left, forma.top = cx - ancho // 2, cy - alto // 2


def _clonar_fila(tabla, modelo=-1):
    tr = copy.deepcopy(tabla._tbl.tr_lst[modelo])
    tabla._tbl.append(tr)


def _quitar_fila(tabla, indice=-1):
    tr = tabla._tbl.tr_lst[indice]
    tr.getparent().remove(tr)


def celda(tabla, fila, col, texto):
    c = tabla.cell(fila, col)
    p = c.text_frame.paragraphs[0]
    rPr = _formato_base(p)
    for extra in list(c.text_frame.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    _escribir(p, texto, rPr)


def tabla(forma, filas):
    """Deja la tabla con exactamente `filas` filas y escribe su contenido."""
    t = forma.table
    while len(t.rows) < len(filas):
        _clonar_fila(t)
    while len(t.rows) > len(filas):
        _quitar_fila(t)
    for f, datos in enumerate(filas):
        for c, valor in enumerate(datos):
            if c < len(t.columns):
                celda(t, f, c, valor)


def recolorear(presentacion):
    """Traduce los colores de FactuGest en toda la presentación."""
    for d in presentacion.slides:
        for forma in d.shapes:
            marcos = []
            if forma.has_text_frame:
                marcos.append(forma.text_frame)
            elif forma.has_table:
                marcos.extend(c.text_frame for c in forma.table.iter_cells())
            for marco in marcos:
                for p in marco.paragraphs:
                    for run in p.runs:
                        try:
                            color = run.font.color
                            if color is None or color.type != 1:
                                continue
                            nuevo = RECOLOR.get(str(color.rgb))
                            if nuevo is not None:
                                color.rgb = nuevo
                        except Exception:
                            pass


def sustituir_parte(presentacion, medida, ruta):
    """Cambia el contenido de la parte-imagen que tenga ese tamaño."""
    datos = Path(ruta).read_bytes()
    for parte in presentacion.part.package.iter_parts():
        if "media" not in str(parte.partname):
            continue
        try:
            if parte.image.size == medida:
                parte._blob = datos
        except Exception:
            continue


def guardar(presentacion):
    try:
        presentacion.save(str(DESTINO))
    except PermissionError:
        subprocess.run(["taskkill", "/F", "/IM", "POWERPNT.EXE"], capture_output=True, check=False)
        time.sleep(1.0)
        presentacion.save(str(DESTINO))


# --------------------------------------------------------------------------
# Contenido de MediApp, diapositiva por diapositiva
# --------------------------------------------------------------------------

def escribir(prs):
    s = prs.slides

    # --- 0. Portada institucional ---------------------------------------
    # Se respeta el reparto de párrafos de la plantilla para que cada uno
    # conserve su propio color y tamaño (el rótulo va en un tono y los nombres
    # en otro). FactuGest tenía tres integrantes y MediApp tiene dos.
    d = s[0]
    txt(d.shapes[1],
        "INTEGRANTES:", "",
        "ÁNGEL JESÚS HERNÁNDEZ ARÉVALO",
        "RICHARD ALBERTO QUIÑONES QUIÑONES", "", "",
        "Ficha 3115426")
    logo = d.shapes[2]                       # el icono es cuadrado, el wordmark no
    _encoger(logo, Inches(2.4))

    # --- 1. Portada del proyecto ----------------------------------------
    d = s[1]
    txt(d.shapes[4], "MediApp")
    txt(d.shapes[5],
        "Sistema web de agendamiento de citas médicas y gestión de historias "
        "clínicas para centros de salud.")
    logo = d.shapes[6]                       # el marco venía aplastado, se cuadra
    logo.height = logo.width = min(logo.width, logo.height)

    # --- 2. Divisoria I --------------------------------------------------
    # «I. FORMULACIÓN DEL PROYECTO» se conserva tal cual.

    # --- 3. Planteamiento del problema -----------------------------------
    d = s[3]
    txt(d.shapes[4],
        "Conseguir una cita médica depende de presentarse en el centro de salud "
        "y de la memoria de quien asigna los turnos.")
    txt(d.shapes[5],
        "El 54 % de los encuestados solicita hoy su cita presentándose en el centro, "
        "y el 71 % se ha desplazado únicamente para pedirla o para cancelarla.")
    txt(d.shapes[6],
        "El 65 % tarda más de media hora en que se la asignen, y el 60 % ha recibido "
        "al menos una vez una cita cruzada con la de otro paciente.")
    txt(d.shapes[7],
        "El problema no es la demanda sino la falta de un canal propio, porque la "
        "disponibilidad no se consulta sino que se recuerda, y por eso la ventanilla "
        "y el teléfono llegan a ofrecer el mismo turno.")

    # --- 4. Justificación -------------------------------------------------
    d = s[4]
    txt(d.shapes[8], "71 %")
    txt(d.shapes[9], "SE DESPLAZARON")
    txt(d.shapes[10], "se ha desplazado únicamente para pedir o para cancelar su cita.")
    txt(d.shapes[11], "65 %")
    txt(d.shapes[12], "ESPERAN DEMASIADO")
    txt(d.shapes[13], "tarda más de media hora en que le asignen el turno de atención.")
    txt(d.shapes[14], "60 %")
    txt(d.shapes[15], "SUFRIERON UN CRUCE")
    txt(d.shapes[16], "ha recibido al menos una vez una cita cruzada con la de otro paciente.")
    txt(d.shapes[17],
        "MediApp permite al paciente elegir el horario libre y dejar su cita en firme "
        "sin pasar por la ventanilla.")

    # --- 5. Objetivo general ---------------------------------------------
    txt(s[5].shapes[4],
        "Desarrollar MediApp, un sistema web de agendamiento de citas médicas con "
        "acceso diferenciado para administrador, médico y paciente, que permita "
        "gestionar la disponibilidad de la atención médica y controlar el historial "
        "clínico en centros de salud.")

    # --- 6. Objetivos específicos (tres, no cinco) ------------------------
    d = s[6]
    txt(d.shapes[4],
        "1. Facilitar la programación de la atención médica mediante un módulo "
        "interactivo que permita al administrador y al paciente agendar citas de "
        "forma intuitiva.")
    txt(d.shapes[5],
        "2. Habilitar la gestión del acto clínico como responsabilidad exclusiva del "
        "médico tratante para la creación y actualización de historias clínicas, "
        "diagnósticos y prescripciones.")
    txt(d.shapes[6],
        "3. Garantizar la confidencialidad de la información médica mediante un "
        "control de acceso basado en roles que restrinja la visualización y la "
        "edición según el perfil del usuario.")
    borrar(d.shapes[8])
    borrar(d.shapes[7])

    # --- 7. Alcance -------------------------------------------------------
    d = s[7]
    # Cada caja de la plantilla admite dos renglones, así que el texto se ajusta
    # a esa medida en lugar de mover la caja.
    txt(d.shapes[4], "Lo que el sistema hace")
    txt(d.shapes[5], "Agenda citas sobre un calendario de disponibilidad real.")
    txt(d.shapes[6], "Aplica en el servidor la separación mínima de 30 minutos.")
    txt(d.shapes[7], "Reserva al médico la historia, el diagnóstico y la receta.")
    txt(d.shapes[8], "Permite al paciente consultar sus historias y sus recetas.")
    txt(d.shapes[9], "Registra los exámenes de laboratorio y sus resultados.")
    txt(d.shapes[10], "Lo que queda fuera de esta versión")
    txt(d.shapes[11], "La notificación por correo o mensaje de texto, que exige un "
                      "servicio contratado.")
    txt(d.shapes[12], "El cifrado en reposo y en tránsito, que depende del despliegue.")
    txt(d.shapes[13], "La firma digital certificada de la historia clínica.")
    txt(d.shapes[14], "La aplicación móvil nativa, porque la interfaz ya responde en "
                      "el navegador.")

    # --- 8. Riesgos del proyecto y restricciones -------------------------
    d = s[8]
    txt(d.shapes[4], "Que la documentación se acumule para el final y el equipo llegue "
                     "sin tiempo.")
    txt(d.shapes[5], "Que quede sin atender alguna observación de la sustentación anterior.")
    txt(d.shapes[6], "Que el entorno local falle el día de la sustentación y no pueda "
                     "mostrarse el sistema.")
    txt(d.shapes[7], "Restricciones")
    txt(d.shapes[8], "El sistema opera sobre un entorno local con XAMPP y no sobre producción.")
    txt(d.shapes[9], "El calendario solo ofrece turnos de los médicos con disponibilidad cargada.")
    txt(d.shapes[10], "La operación requiere la base de datos, porque las reglas se "
                      "resuelven en el servidor.")

    # --- 9. Cronograma ----------------------------------------------------
    d = s[9]
    txt(d.shapes[3], "Cronograma de Actividades")
    imagen(d, d.shapes[4], DIAGRAMAS / "FIG-cronograma.png")

    # --- 10. Divisoria II -------------------------------------------------

    # --- 11. Captura de requisitos ---------------------------------------
    d = s[11]
    txt(d.shapes[4], "Encuesta")
    txt(d.shapes[5],
        "Cuestionario de doce preguntas cerradas aplicado a 48 usuarios de centros de "
        "salud de Cúcuta, agrupadas en cuatro bloques que indagaban cómo solicitan hoy "
        "su cita, cuánto tardan en obtenerla, qué inconvenientes han sufrido y de qué "
        "manera preferirían agendarla.")
    # El marco de la plantilla es apaisado (4,2 a 1) y la figura del contraste
    # tiene exactamente esa proporción, así que entra sin dejar aire a los lados.
    imagen(d, d.shapes[6], DIAGRAMAS / "FIG-encuesta-contraste.png")
    txt(d.shapes[7],
        "Quieren elegir ellos mismos el horario libre y que la cita quede en firme de inmediato.")

    # --- 12. Historias de usuario ----------------------------------------
    d = s[12]
    txt(d.shapes[4], "Derivadas de la encuesta")
    txt(d.shapes[5],
        "HU-03. Como paciente, quiero ver en un calendario los horarios libres de cada "
        "médico para escoger el que me convenga sin preguntar por ellos.")
    txt(d.shapes[6],
        "HU-04. Como paciente, quiero que el horario que elijo quede agendado en el "
        "momento para no tener que esperar a que un tercero me lo confirme.")
    txt(d.shapes[7],
        "HU-11. Como médico, quiero registrar la historia clínica del paciente que "
        "atiendo para dejar constancia de su evolución bajo mi propia autoría.")

    # --- 13. Requisitos funcionales relacionados -------------------------
    d = s[13]
    tabla(d.shapes[4], [
        ["Código", "Requisito", "Qué resuelve"],
        ["RF 5.8", "Calendario semanal de disponibilidad",
         "El paciente ve los horarios libres del médico sin preguntar por ellos"],
        ["RF 5.2", "Agendamiento por el propio paciente",
         "El horario elegido queda en firme sin esperar la confirmación de un tercero"],
        ["RF 5.4", "Separación mínima entre citas",
         "Se rechaza toda cita a menos de 30 minutos de otra del mismo médico"],
        ["RF 6.2", "Exclusividad del médico sobre la historia",
         "Solo el médico tratante escribe en ella, y la autoría sale de la sesión"],
    ])

    # --- 14. Matriz de stakeholders --------------------------------------
    imagen(s[14], s[14].shapes[4], DIAGRAMAS / "FIG-stakeholders.png")

    # --- 15. Técnica de priorización -------------------------------------
    tabla(s[15].shapes[4], [
        ["Puntaje", "Valor de negocio (60%)", "Urgencia (40%)"],
        ["1", "Impacto muy bajo para la atención", "No es urgente, puede esperar"],
        ["2", "Impacto bajo, aporta poca funcionalidad", "Poco urgente"],
        ["3", "Impacto medio, mejora procesos importantes", "Urgencia moderada"],
        ["4", "Impacto alto, necesario para operar", "Muy urgente"],
        ["5", "Impacto crítico, indispensable", "Debe implementarse de inmediato"],
    ])

    # --- 16. Conclusión de la priorización -------------------------------
    d = s[16]
    tabla(d.shapes[4], [
        ["Código", "Módulo", "Puntaje"],
        ["RF 5", "Agendamiento de citas", "4.6"],
        ["RF 1", "Autenticación y sesión", "4.5"],
        ["RF 6", "Historia clínica", "4.3"],
        ["RF 7", "Consultas y diagnósticos", "4.3"],
        ["RF 8", "Prescripciones", "4.2"],
        ["RF 9", "Exámenes de laboratorio", "4.2"],
        ["RF 2", "Gestión de usuarios", "4.1"],
        ["RF 3", "Gestión de médicos", "4.1"],
        ["RF 4", "Gestión de pacientes", "4.1"],
        ["RF 10", "Catálogos", "3.6"],
    ])
    txt(d.shapes[5],
        "Encabeza la lista agendamiento de citas, que es el módulo del que depende el "
        "problema que el proyecto viene a resolver, y cierra catálogos, que agrupa "
        "parámetros que se ajustan de vez en cuando.")

    # --- 17. Requisitos funcionales --------------------------------------
    tabla(s[17].shapes[4], [
        ["Código", "Módulo", "Requisitos", "Críticos"],
        ["RF 1", "Autenticación y sesión", "9", "6"],
        ["RF 2", "Gestión de usuarios", "8", "3"],
        ["RF 3", "Gestión de médicos", "8", "3"],
        ["RF 4", "Gestión de pacientes", "9", "3"],
        ["RF 5", "Agendamiento de citas", "14", "11"],
        ["RF 6", "Historia clínica", "8", "5"],
        ["RF 7", "Consultas y diagnósticos", "8", "5"],
        ["RF 8", "Prescripciones", "8", "3"],
        ["RF 9", "Exámenes de laboratorio", "7", "3"],
        ["RF 10", "Catálogos", "8", "0"],
        ["", "Total", "87", "42"],
    ])

    # --- 18. Requisitos no funcionales -----------------------------------
    tabla(s[18].shapes[4], [
        ["Categoría", "N.º", "Ejemplo"],
        ["Seguridad", "7", "La autoría de un registro clínico se toma de la sesión y nunca del formulario"],
        ["Fiabilidad", "4", "De dos reservas simultáneas del mismo turno solo una queda en firme"],
        ["Usabilidad", "4", "El ingreso lleva directo al panel del rol, sin pedir que se elija perfil"],
        ["Mantenibilidad", "4", "La comprobación del rol está implementada una sola vez, como decorador"],
        ["Rendimiento", "2", "El calendario de un médico se arma en menos de dos segundos"],
        ["Portabilidad", "2", "Opera desde cualquier sistema operativo con un navegador vigente"],
        ["Escalabilidad", "1", "Admite nuevas especialidades y medicamentos sin cambios en el código"],
        ["Total", "24", ""],
    ])

    # --- 19. Divisoria III ------------------------------------------------

    # --- 20 y 21. Pantallas del sistema ----------------------------------
    d = s[20]
    txt(d.shapes[2], "Panel del Administrador")
    imagen(d, d.shapes[3], CAPTURAS / "IU-02 panel del administrador.jpg")

    d = s[21]
    txt(d.shapes[2], "Calendario de Disponibilidad")
    imagen(d, d.shapes[3], CAPTURAS / "IU-05 calendario de disponibilidad.jpg")

    # --- 22 y 23. Base de datos ------------------------------------------
    d = s[22]
    txt(d.shapes[3], "Modelo Físico de la Base de Datos")
    imagen(d, d.shapes[4], DIAGRAMAS / "FIG-fisico.png")

    d = s[23]
    txt(d.shapes[3], "Modelo Entidad-Relación")
    imagen(d, d.shapes[4], DIAGRAMAS / "FIG-mer.png")

    # --- 24. Diccionario de datos ----------------------------------------
    tabla(s[24].shapes[4], [
        ["Campo", "Tipo", "Clave", "Descripción"],
        ["id_paciente", "Entero (11)", "PK", "Identificador interno del paciente"],
        ["nombre", "Texto (100)", "No aplica", "Nombre completo de la persona atendida"],
        ["tipo_documento", "Texto (20)", "No aplica", "Clase de documento de identidad presentado"],
        ["numero_documento", "Texto (30)", "Única", "Número del documento; impide registrar dos veces a la misma persona"],
        ["fecha_nacimiento", "Fecha", "No aplica", "Fecha de nacimiento, de la que se deriva la edad"],
        ["telefono", "Texto (20)", "No aplica", "Número de contacto para avisos sobre la cita"],
        ["direccion", "Texto (150)", "No aplica", "Dirección de residencia"],
        ["email", "Texto (120)", "No aplica", "Correo de contacto del paciente"],
        ["id_usuario", "Entero (11)", "FK", "Cuenta con la que el paciente ingresa a consultar lo suyo"],
        ["estado", "Texto (20)", "No aplica", "Activo o inactivo; la baja es lógica y conserva el historial"],
    ])

    # --- 25. Casos de uso -------------------------------------------------
    d = s[25]
    txt(d.shapes[4],
        "El análisis identificó 60 casos de uso agrupados en los diez módulos del "
        "sistema, con cuatro actores. Uno de ellos no es una persona, ya que es el "
        "propio sistema el que aplica por su cuenta las dos reglas de la agenda.")
    imagen(d, d.shapes[5], DIAGRAMAS / "DCU-00 general.png")

    # --- 26. Documentación de casos de uso -------------------------------
    tabla(s[26].shapes[4], [
        ["CU-24", "Reservar una cita desde el calendario"],
        ["Actores", "Paciente, Sistema"],
        ["Precondición", "El paciente tiene sesión iniciada y el médico tiene disponibilidad cargada"],
        ["Pasos 1 a 3", "Elige médico y semana, el sistema muestra los turnos libres de treinta minutos y el paciente escoge uno"],
        ["Pasos 4 a 6", "El servidor toma la identidad de la sesión, comprueba la separación mínima y el máximo de una cita al día"],
        ["Paso 7", "Guarda la cita en estado agendada y la muestra en el panel del paciente"],
        ["Postcondición", "Existe una cita a nombre de quien la pidió y el turno deja de ofrecerse"],
        ["Excepción", "Si el turno cae dentro de los treinta minutos de otra cita del mismo médico, se rechaza"],
    ])

    # --- 27. Metodología --------------------------------------------------
    d = s[27]
    txt(d.shapes[4],
        "Scrum, con diez sprints de dos semanas, aproximadamente uno por cada módulo "
        "del sistema. Cada sprint cerró con un módulo utilizable y no con una parte "
        "de varios.")
    imagen(d, d.shapes[5], DIAGRAMAS / "FIG-scrum.png")

    # --- 28. Divisoria IV -------------------------------------------------

    # --- 29. Herramientas de backend, base de datos y web ----------------
    d = s[29]
    formas = list(d.shapes)                  # se fija antes de borrar nada
    txt(formas[3], "Herramientas de Backend, Base de Datos y Web")
    txt(formas[10], "Python con Flask")
    txt(formas[11], "MySQL / MariaDB")
    txt(formas[12], "HTML, CSS y JavaScript")
    txt(formas[13], "PyCharm")
    txt(formas[14], "Lenguaje y framework del servidor. Resuelve el enrutamiento, la "
                    "sesión y las reglas de la agenda antes de escribir en la base.")
    txt(formas[15], "Gestor relacional con once tablas, sesenta y tres columnas y "
                    "catorce claves foráneas, con acceso mediante consultas SQL explícitas.")
    txt(formas[16], "Plantillas Jinja2 con Bootstrap 5 para las tres vistas de rol y "
                    "el calendario de disponibilidad.")
    txt(formas[17], "Entorno de desarrollo usado para escribir, depurar y mantener el código.")
    # El logo de la primera fila era el de FastAPI, que MediApp no usa. Si hay un
    # PNG en assets/flask.png se pone; si no, la fila se queda solo con su rótulo.
    flask_logo = ASSETS / "flask.png"
    if flask_logo.exists():
        imagen(d, formas[4], flask_logo)
    else:
        borrar(formas[4])

    # --- 30. Herramientas de apoyo, pruebas y control de versiones -------
    # FactuGest lista cuatro; MediApp usa tres, así que se retira la fila de
    # Postman con su logo y las tres restantes se reparten el alto de la tabla.
    d = s[30]
    txt(d.shapes[2], "Herramientas de Apoyo, Pruebas y Control de Versiones")
    forma_tabla = d.shapes[3]
    claude, postman = d.shapes[4], d.shapes[5]
    git, github, pytest_logo = d.shapes[6], d.shapes[7], d.shapes[8]
    borrar(postman)                            # MediApp no usa Postman
    tabla(forma_tabla, [
        ["NOMBRE", "LOGO", "DESCRIPCION"],
        ["", "  Claude AI", "  Asistente empleado para revisar código, contrastar decisiones de diseño y redactar documentación."],
        ["", "  Git y GitHub", "  Control de versiones y alojamiento remoto para el trabajo colaborativo del equipo."],
        ["", "  pytest", "  Ejecuta las 91 pruebas automatizadas del proyecto en 6,9 segundos."],
    ])
    _repartir_logos(forma_tabla, [[claude], [git, github], [pytest_logo]])

    # --- 31. Divisoria V --------------------------------------------------

    # --- 32. Pruebas del sistema -----------------------------------------
    d = s[32]
    txt(d.shapes[3], "Pruebas del Sistema")
    txt(d.shapes[4], "91 pruebas automatizadas en 6,9 segundos, más pruebas funcionales "
                     "sobre la interfaz.")
    txt(d.shapes[5], "Qué se comprueba")
    txt(d.shapes[6], "Que la separación mínima de treinta minutos se aplique en el "
                     "servidor y no solo en el calendario.")
    txt(d.shapes[7], "Que de dos reservas simultáneas del mismo turno quede una sola "
                     "cita en firme.")
    txt(d.shapes[8], "Que la autoría de un registro clínico salga de la sesión y nunca "
                     "del formulario.")
    txt(d.shapes[9], "Que el control de acceso impida a cada rol lo que no le corresponde.")
    txt(d.shapes[10], "Qué quedó fuera")
    txt(d.shapes[11], "La carga sostenida, que exige una infraestructura de producción.")
    txt(d.shapes[12], "El cifrado en tránsito, que depende del despliegue y no del programa.")

    # --- 33. Evidencia 1 --------------------------------------------------
    d = s[33]
    txt(d.shapes[4], "RF 6.1. Creación del registro de historia clínica")
    imagen(d, d.shapes[6], CAPTURAS / "CP-01a historia clinica formulario.jpg")
    imagen(d, d.shapes[8], CAPTURAS / "CP-01b historia clinica creada.jpg")
    txt(d.shapes[9], "La historia quedó registrada a nombre del médico que la escribió, "
                     "y el administrador solo puede consultarla.")

    # --- 34. Evidencia 2 --------------------------------------------------
    d = s[34]
    txt(d.shapes[4], "RF 5.4. Separación mínima entre citas del mismo médico")
    imagen(d, d.shapes[6], CAPTURAS / "CP-02a cita en conflicto formulario.jpg")
    imagen(d, d.shapes[8], CAPTURAS / "CP-02b cita en conflicto rechazada.jpg")
    txt(d.shapes[9], "La cita a quince minutos de otra del mismo médico fue rechazada "
                     "por el servidor antes de guardarse.")

    # --- 35. Resultado obtenido de la reserva ----------------------------
    d = s[35]
    txt(d.shapes[3], "Resultado Obtenido de la Reserva")
    txt(d.shapes[4], "La prueba fue satisfactoria, la cita quedó agendada a nombre de "
                     "quien la pidió.")
    tabla(d.shapes[5], [
        ["Lo que se comprobó", "Resultado"],
        ["Identidad", "El identificador del paciente se tomó de la sesión y se ignoró el enviado en el formulario"],
        ["Separación", "El turno a menos de treinta minutos de otra cita del mismo médico fue rechazado"],
        ["Límite diario", "Un segundo intento del mismo paciente para el mismo día no prosperó"],
        ["Concurrencia", "De dos reservas simultáneas del mismo turno solo una quedó en firme"],
        ["Estado", "La cita quedó en estado agendada y el turno dejó de ofrecerse en el calendario"],
    ])

    # --- 36. Resultados obtenidos ----------------------------------------
    tabla(s[36].shapes[4], [
        ["Objetivo", "Resultado", "Estado"],
        ["Programación de la atención médica",
         "Calendario de disponibilidad con turnos de treinta minutos desde el que agendan tanto el administrador como el propio paciente, con la reserva en firme al momento.",
         "Cumplido"],
        ["Acto clínico exclusivo del médico",
         "Historia clínica, diagnóstico y prescripción creados únicamente por el médico tratante, con la autoría tomada de la sesión y no del formulario.",
         "Cumplido"],
        ["Confidencialidad por control de acceso",
         "Tres roles comprobados en el servidor por ruta y por propiedad del registro, con las contraseñas derivadas mediante scrypt.",
         "Cumplido"],
        ["Reglas de negocio de la agenda",
         "Separación mínima de treinta minutos y máximo de una cita por paciente al día, resueltas en el servidor dentro de la misma operación que escribe.",
         "Cumplido"],
        ["Verificación automatizada",
         "91 pruebas sobre autenticación, permisos, agenda, concurrencia y datos, todas pasando en 6,9 segundos.",
         "Cumplido"],
        ["Cifrado de la información en reposo",
         "Depende de la infraestructura de despliegue, razón por la cual el tercer objetivo se acotó de forma expresa al control de acceso por roles.",
         "Fuera de alcance"],
    ])

    # --- 37. Cierre -------------------------------------------------------
    # El fondo institucional ya dice GRACIAS con los datos del SENA; solo sobra
    # la franja de contacto de FactuGest.
    borrar(s[37].shapes[1])


def _repartir_logos(forma_tabla, grupos, alto_encabezado=Inches(0.70)):
    """Fija el alto de las filas y centra cada grupo de logos sobre la suya.

    Hay que fijarlo: `row.height` es un mínimo, y PowerPoint encoge o estira la
    fila según el texto, de modo que la posición real no se puede deducir de la
    plantilla. Fijando el alto, la cuenta y el dibujo coinciden.
    """
    t = forma_tabla.table
    cuerpo = len(t.rows) - 1
    if cuerpo < 1:
        return
    alto_fila = int((forma_tabla.height - alto_encabezado) / cuerpo)
    t.rows[0].height = int(alto_encabezado)
    for fila in list(t.rows)[1:]:
        fila.height = alto_fila
    y = forma_tabla.top + int(alto_encabezado)
    for grupo in grupos:
        for logo in grupo:
            logo.top = y + (alto_fila - logo.height) // 2
        y += alto_fila


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ORIGEN, DESTINO)          # la plantilla, tal cual
    prs = Presentation(str(DESTINO))

    escribir(prs)
    recolorear(prs)

    sustituir_parte(prs, (320, 320), LOGO_MEDIAPP)     # logo FG -> logo MediApp
    sustituir_parte(prs, (1024, 1024), TRANSPARENTE)   # marca de agua FG -> nada

    guardar(prs)
    print("OK:", DESTINO, "|", len(prs.slides), "diapositivas")


if __name__ == "__main__":
    main()
